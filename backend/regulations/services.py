"""Bounded text extraction; uploads remain private."""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, BadZipFile

from django.core.exceptions import ValidationError

MAX_UPLOAD = 5 * 1024 * 1024
MAX_TEXT = 100000


def extract_text(upload):
    if upload.size > MAX_UPLOAD:
        raise ValidationError("Tệp tối đa 5 MB.")
    suffix = Path(upload.name).suffix.lower()
    raw = upload.read(MAX_UPLOAD + 1)
    upload.seek(0)
    if len(raw) > MAX_UPLOAD:
        raise ValidationError("Tệp tối đa 5 MB.")
    try:
        if suffix in {".txt", ".md"}:
            text = raw.decode("utf-8-sig")
        elif suffix == ".docx":
            # Reject compressed archives with excessive expanded size.
            with ZipFile(BytesIO(raw)) as archive:
                if (
                    sum(item.file_size for item in archive.infolist())
                    > 20 * 1024 * 1024
                ):
                    raise ValidationError(
                        "Tệp Word có kích thước giải nén quá lớn."
                    )
            from docx import Document

            doc = Document(BytesIO(raw))
            chunks = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                chunks.extend(
                    " | ".join(cell.text for cell in row.cells)
                    for row in table.rows
                )
            text = "\n".join(chunks)
        elif suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(raw))
            if reader.is_encrypted or len(reader.pages) > 80:
                raise ValidationError(
                    "PDF phải không có mật khẩu và tối đa 80 trang."
                )
            chunks = []
            count = 0
            for page in reader.pages:
                content = page.extract_text() or ""
                count += len(content)
                if count > MAX_TEXT:
                    raise ValidationError("Nội dung tối đa 100.000 ký tự.")
                chunks.append(content)
            text = "\n".join(chunks)
        else:
            raise ValidationError(
                "Chỉ chấp nhận TXT, MD, DOCX hoặc PDF có văn bản."
            )
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError(
            (
                "Không đọc được tệp. Kiểm tra định dạ"
                "ng và dùng UTF-8 cho tệp văn bản."
            )
        ) from exc
    text = text.strip()
    if not text:
        raise ValidationError(
            "Không tìm thấy văn bản. PDF ảnh quét cần OCR trước khi tải lên."
        )
    if len(text) > MAX_TEXT:
        raise ValidationError("Nội dung tối đa 100.000 ký tự.")
    return text
